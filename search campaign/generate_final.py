import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment
from copy import copy

input_file  = "Search campaign antigravity.xlsx"
output_file = "ditevrozvodu_Ads_Final.xlsx"

# ---------------------------------------------------------------------------
# AD COPY — all 8 ad groups (validated, all within limits)
# ---------------------------------------------------------------------------
ads = [
    # 1.1  Jak mluvit s dětmi o rozvodu
    dict(
        Campaign="Rozvod a děti",
        Ad_group="Jak mluvit s dětmi o rozvodu",
        h1="Jak říct dětem o rozvodu?",
        h2="Najděte ta správná slova",
        h3="Jak oznámit rozvod bez slz?",
        h4="Nerozbijte dětem svět",
        h5="Rozvod bez dětských traumat",
        h6="Tápete, jak jim to říct?",
        h7="Řekněte jim to citlivě",
        h8="Slova, která dětem pomohou",
        h9="Průvodce těžkým rozhovorem",
        h10="Ochrana dětí na 1. místě",
        d1="Hledáte správná slova? Poradíme vám, jak s dětmi mluvit otevřeně a přitom je ochránit.",
        d2="Rozvod bolí, ale nemusí zranit. Zjistěte, jak oznámit konec vztahu bez dětských traumat.",
        d3="Bojíte se dětských slz? Naučte se vést ten nejtěžší rozhovor s maximálním citem a klidem.",
        d4="Vaše děti si zaslouží pravdu i bezpečí. Praktické rady od zkušených psychologů zdarma.",
        p1="deti", p2="komunikace",
        url="https://ditevrozvodu.cz/jak-mluvit-s-detmi-o-rozvodu",
    ),
    # 1.2  Děti a rozpad rodiny
    dict(
        Campaign="Rozvod a děti",
        Ad_group="Děti a rozpad rodiny",
        h1="Rozvod rodičů, ne dětí",
        h2="Jak rozvod ovlivňuje děti?",
        h3="Zachraňte úsměv svých dětí",
        h4="Konec vztahu, ne rodiny",
        h5="Jak ochránit děti v rozvodu",
        h6="Rozpad rodiny bez paniky",
        h7="Vaše děti rozvod zvládnou",
        h8="Zůstaňte skvělými rodiči",
        h9="Dopad rozvodu na děti: Rady",
        h10="Ochranný štít pro vaše děti",
        d1="Manželství končí, vy zůstáváte rodiči. Zjistěte, jak dětem v rozvodu zachovat domov.",
        d2="Strach o děti? Poradíme, jak jim dodat jistotu i ve chvílích, kdy se všechno mění.",
        d3="Rozvod rodičů zásadně mění dětský svět. Přečtěte si, jak ho co nejvíce ochránit.",
        d4="Udělejte to jinak. Rozveďte se tak, aby vaše děti netrpěly. Jde to i bez boje.",
        p1="deti", p2="rozpad-rodiny",
        url="https://ditevrozvodu.cz/deti-a-rozpad-rodiny",
    ),
    # 1.3  Psychologie dítěte při rozvodu
    dict(
        Campaign="Rozvod a děti",
        Ad_group="Psychologie dítěte při rozvodu",
        h1="Co prožívá vaše dítě?",
        h2="Emoce dítěte při rozvodu",
        h3="Zachraňte dětskou psychiku",
        h4="Jak pomoci dítěti v rozvodu",
        h5="Kdy zasáhnout? Rady psychologa",
        h6="Zbavte dítě pocitu viny",
        h7="Jak dítě zvládne rozvod rodičů",
        h8="Dítě na prvním místě",
        h9="Psycholog dítěte a rozvod",
        h10="Rozvod bez dětské jizvy",
        d1="Děti prožívají rozvod jinak než dospělí. Nahlédněte do jejich světa a dejte jim oporu.",
        d2="Strach, vztek i pocit viny. Pochopte, čím si dítě prochází, a buďte mu bezpečným břehem.",
        d3="Trápí se vaše dítě? Rady od psychologů vám ukážou, jak číst emoce a včas zasáhnout.",
        d4="Nečekejte, až se problémy vyhrotí. Poskytněte dítěti psychologickou pomoc, co si zaslouží.",
        p1="deti", p2="psychologie",
        url="https://ditevrozvodu.cz/psychologie-ditete-pri-rozvodu",
    ),
    # 2.1  Online kurz rozvod děti
    dict(
        Campaign="Online kurz - rozvod",
        Ad_group="Online kurz rozvod děti",
        h1="Online kurz: Rozvod a děti",
        h2="Kurz pro rozvádějící se rodiče",
        h3="Zvládněte rozvod z domova",
        h4="Tajemství klidného rozvodu",
        h5="Video lekce pro náročné chvíle",
        h6="Krok za krokem těžkým obdobím",
        h7="Kurz, co chrání vaše děti",
        h8="Rozvod chytře a bez traumat",
        h9="Naučte se mluvit s dětmi",
        h10="Investujte do klidu rodiny",
        d1="Žádné dojíždění, jen čistá praxe. Kurz vám dá ověřené nástroje pro klidný rozvod s dětmi.",
        d2="Ztrácíte půdu pod nohama? Online kurz vás provede rozvodem s dětmi od A do Z.",
        d3="Srozumitelné lekce od psychologů. Přidejte se k rodičům, kteří to zvládli lépe, čekali.",
        d4="Konec bezesných nocí. Praktická videa vám dají jistotu v každém kroku. Zkuste to dnes.",
        p1="kurz", p2="deti",
        url="https://ditevrozvodu.cz/online-kurz",
    ),
    # 2.2  Pomoc při rozvodu
    dict(
        Campaign="Online kurz - rozvod",
        Ad_group="Pomoc při rozvodu",
        h1="Nezvládáte to? Pomůžeme vám",
        h2="S rozvodem na to nejste sami",
        h3="Jak přežít rozvod jako rodič",
        h4="Okamžitá podpora při rozvodu",
        h5="Krizová pomoc pro rodiče",
        h6="Rady pro rodiče v rozvodu",
        h7="Únava z rozvodu? Pomůžeme",
        h8="Průvodce pro rodiče v tísni",
        h9="Jak zvládnout rozvod? Víme.",
        h10="Opora, o niž se lze opřít",
        d1="Stojíte na pokraji sil? Nechte si pomoct. Společně najdeme cestu z rozvodové krize.",
        d2="Hádky, únava, strach o děti. Obraťte se na profesionály, kteří vědí, jak uklidnit bouři.",
        d3="Vaše problémy mají řešení. Získejte klid pro sebe i bezpečí pro vaše děti. Jsme tu.",
        d4="Poradíme vám, jak zvládnout rozvod jako skvělý rodič. Ověřené rady z praxe čekají na vás.",
        p1="pomoc", p2="rozvod",
        url="https://ditevrozvodu.cz/pomoc-pri-rozvodu",
    ),
    # 3.1  Péče o děti po rozvodu
    dict(
        Campaign="Praktické info - rozvod",
        Ad_group="Péče o děti po rozvodu",
        h1="Péče o děti po rozvodu",
        h2="Střídavá, nebo výlučná péče?",
        h3="Svěření do péče: Jak na to",
        h4="Dohoda o dětech bez hádek",
        h5="Zákon a péče o děti 2026",
        h6="Soud, nebo rozumná dohoda?",
        h7="Co je střídavá péče v praxi",
        h8="Nejlepší péče pro vaše dítě",
        h9="Domluvte se jako dospělí",
        h10="Práva rodičů v rozvodu",
        d1="Střídavá, výlučná, nebo společná péče? Zjistěte, který model dá vašim dětem stabilitu.",
        d2="Nová pravidla 2026: Co se mění v péči o děti po rozvodu? Čtěte přehledný průvodce.",
        d3="Nehádejte se o děti. Poradíme vám, jak nastavit spravedlivá pravidla ku prospěchu všech.",
        d4="Soud nerozhodne lépe než vy. Získejte návod, jak sepsat fungující rodičovskou dohodu.",
        p1="pece", p2="deti",
        url="https://ditevrozvodu.cz/pece-o-deti-po-rozvodu",
    ),
    # 3.2  Výživné a soud
    dict(
        Campaign="Praktické info - rozvod",
        Ad_group="Výživné a soud",
        h1="Výživné na dítě: Jak na to",
        h2="Kolik dostanete na dítě?",
        h3="Jak probíhá rozvod u soudu?",
        h4="Konec dohadů o alimentech",
        h5="Kalkulačka výše výživného",
        h6="OSPOD a rozvod: Co vás čeká",
        h7="Soudní řízení s dětmi jasně",
        h8="Rozvod s dětmi: Kolik to stojí",
        h9="Připravte se na rozvod u soudu",
        h10="Férové výživné bez překvapení",
        d1="Peníze jsou hlavní kámen úrazu. Spočítejte si férové výživné a ušetřete nervy i čas.",
        d2="Jak určí alimenty soud? Vyzbrojte se informacemi dřív, než přijde nepříjemné překvapení.",
        d3="Co obnáší rozvod s dětmi u soudu? Přehled OSPOD, výživného i soudního řízení na místě.",
        d4="Soudní tahanice nikomu neprospějí. Zjistěte, jak vyřešit výživné efektivně a v klidu.",
        p1="vyzivne", p2="soud",
        url="https://ditevrozvodu.cz/vyzivne-a-soud",
    ),
    # 3.3  Mediace a domluva
    dict(
        Campaign="Praktické info - rozvod",
        Ad_group="Mediace a domluva",
        h1="Rozvod dohodou. Jde to!",
        h2="Mediace místo soudního boje",
        h3="Dohoda bez soudu: Jak na to",
        h4="Vyhněte se válce právníků",
        h5="Jak se domluvit s ex o dětech",
        h6="Rodinný mediátor pomůže",
        h7="Komunikace po rozvodu funguje",
        h8="Kompromis místo tvrdého boje",
        h9="Ušetřete soudní tahanice",
        h10="Dohoda šetří čas i peníze",
        d1="Bojíte se, že vás rozvod zruinuje? Mediace je rychlejší, levnější a pro děti lepší.",
        d2="I zarytí nepřátelé dokážou najít smír. Zkuste mimosoudní cestu, ze které profitují děti.",
        d3="Komunikace s ex na bodu mrazu? Mediátor vám pomůže prolomit ledy a najít fér řešení.",
        d4="Rodinná mediace chrání děti od soudního stresu. Zjistěte, jak to celé funguje.",
        p1="mediace", p2="dohoda",
        url="https://ditevrozvodu.cz/mediace-a-domluva",
    ),
]

# ---------------------------------------------------------------------------
# Validate (fail-fast)
# ---------------------------------------------------------------------------
print("=== VALIDATION ===")
all_ok = True
for ad in ads:
    ag = ad["Ad_group"]
    for f in [f"h{i}" for i in range(1, 11)]:
        if len(ad[f]) > 30:
            print(f"  ❌ H TOO LONG ({len(ad[f])}/30) [{ag}] {f}: '{ad[f]}'")
            all_ok = False
    for f in [f"d{i}" for i in range(1, 5)]:
        if len(ad[f]) > 90:
            print(f"  ❌ D TOO LONG ({len(ad[f])}/90) [{ag}] {f}: '{ad[f]}'")
            all_ok = False
    for f in ["p1", "p2"]:
        if len(ad[f]) > 15:
            print(f"  ❌ PATH TOO LONG ({len(ad[f])}/15) [{ag}] {f}: '{ad[f]}'")
            all_ok = False
if all_ok:
    print("  ✅ All lengths OK!")

# ---------------------------------------------------------------------------
# Open the template workbook and clone its structure
# ---------------------------------------------------------------------------
wb = load_workbook(input_file)

# ---- Sheet 1: CZ KWs — leave untouched (already has correct data) ----

# ---- Sheet 2: CZ reklamy — write ad copy, preserving formulas & styles ----
ws = wb["CZ reklamy"]

# Column letter mapping (from the template inspection)
# A=Campaign  B=Ad group
# C=H1  D=LEN  E=H2  F=LEN  G=H3  H=LEN  I=H4  J=LEN  K=H5  L=LEN
# M=H6  N=LEN  O=H7  P=LEN  Q=H8  R=LEN  S=H9  T=LEN  U=H10 V=LEN
# W=D1  X=LEN  Y=D2  Z=LEN  AA=D3 AB=LEN AC=D4  AD=LEN
# AE=P1 AF=LEN AG=P2  AH=LEN AI=Final URL

headline_cols = ["C","E","G","I","K","M","O","Q","S","U"]   # H1–H10
desc_cols     = ["W","Y","AA","AC"]                          # D1–D4
path_cols     = ["AE","AG"]                                  # P1, P2
url_col       = "AI"

for row_idx, ad in enumerate(ads, start=2):
    headlines  = [ad[f"h{i}"] for i in range(1, 11)]
    descs      = [ad[f"d{i}"] for i in range(1, 5)]
    paths      = [ad["p1"], ad["p2"]]

    ws[f"A{row_idx}"] = ad["Campaign"]
    ws[f"B{row_idx}"] = ad["Ad_group"]

    for col, val in zip(headline_cols, headlines):
        ws[f"{col}{row_idx}"] = val

    for col, val in zip(desc_cols, descs):
        ws[f"{col}{row_idx}"] = val

    for col, val in zip(path_cols, paths):
        ws[f"{col}{row_idx}"] = val

    ws[f"{url_col}{row_idx}"] = ad["url"]

    # The LEN formula columns are already present in the template as =LEN(Xn).
    # openpyxl preserves them from the template rows — no need to re-write.

wb.save(output_file)
print(f"\n✅ Saved: {output_file}")
