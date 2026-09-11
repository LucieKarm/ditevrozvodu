import pandas as pd
from openpyxl import load_workbook

wb = load_workbook("Search campaign antigravity.xlsx")

for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    print(f"\n=== Sheet: {sheet_name} ===")
    print(f"Dimensions: {ws.dimensions}")
    print(f"Max row: {ws.max_row}, Max col: {ws.max_column}")
    print("\n--- Row 1 (headers) cell by cell ---")
    for cell in ws[1]:
        if cell.value is not None:
            print(f"  {cell.coordinate}: value='{cell.value}' font_bold={cell.font.bold if cell.font else None} fill={cell.fill.fgColor.rgb if cell.fill and cell.fill.fgColor else None} col_width={ws.column_dimensions[cell.column_letter].width}")
    print("\n--- All rows ---")
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None:
                print(f"  {cell.coordinate}: '{cell.value}'")
